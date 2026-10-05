"""Create and append scan findings to a Word report."""
from __future__ import annotations

from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt

REPORT_PATH = Path(__file__).parent / "reports" / "spring_week_networking_report.docx"


def append_scan(opportunities: list[dict], insights_by_key: dict[str, list[dict]], scanned_at: str | None = None, path: Path = REPORT_PATH) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        doc = Document(path)
    else:
        doc = Document()
        doc.add_heading("Spring Week Finder & Networking Report", 0)
        doc.add_paragraph("Research sourced from publicly accessible web pages. Verify current details on employer pages.")
    stamp = scanned_at or datetime.now().astimezone().strftime("%Y-%m-%d %H:%M %Z")
    doc.add_heading(f"Scan — {stamp}", level=1)
    for opportunity in opportunities:
        doc.add_heading(f"{opportunity.get('company_name') or 'Company'} — {opportunity.get('program_title') or 'Program'}", level=2)
        for label, field in (("Sector", "sector"), ("Application deadline", "application_deadline"), ("Location", "location")):
            doc.add_paragraph(f"{label}: {opportunity.get(field) or 'Not listed'}")
        url = opportunity.get("direct_application_url", "")
        p = doc.add_paragraph("Application: ")
        if url:
            p.add_run(url)
        matches = insights_by_key.get(opportunity.get("direct_application_url", ""), [])
        if matches:
            doc.add_heading("Networking & skill insights", level=3)
            for person in matches:
                doc.add_paragraph(f"{person.get('person_name') or 'Public profile'} — {person.get('current_or_past_role') or 'Role not listed'}")
                doc.add_paragraph(f"Education: {person.get('education') or 'Not listed'}")
                doc.add_paragraph(f"Technical skills: {', '.join(person.get('technical_skills', [])) or 'Not listed'}")
                doc.add_paragraph(f"Soft skills: {', '.join(person.get('soft_skills', [])) or 'Not listed'}")
                if person.get("profile_url"):
                    doc.add_paragraph(f"Profile: {person['profile_url']}")
                if person.get("evidence_note"):
                    doc.add_paragraph(f"Evidence: {person['evidence_note']}")
    doc.save(path)
    return path
